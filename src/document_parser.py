import io
from dataclasses import dataclass
from typing import List

import pdfplumber
from docx import Document
from openpyxl import load_workbook


@dataclass
class ParsedDocument:
    name: str
    text: str


def parse_document(name: str, content: bytes) -> ParsedDocument:
    lower_name = name.lower()
    if lower_name.endswith(".pdf"):
        return ParsedDocument(name=name, text=_parse_pdf(content))
    if lower_name.endswith(".docx"):
        return ParsedDocument(name=name, text=_parse_docx(content))
    if lower_name.endswith(".xlsx"):
        return ParsedDocument(name=name, text=_parse_xlsx(content))
    if lower_name.endswith(".txt"):
        return ParsedDocument(name=name, text=content.decode("utf-8", errors="ignore"))
    return ParsedDocument(name=name, text="")


def _parse_pdf(content: bytes) -> str:
    text_parts: List[str] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                text_parts.append(page_text)
    return "\n".join(text_parts)


def _parse_docx(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    return "\n".join([para.text for para in doc.paragraphs if para.text.strip()])


def _parse_xlsx(content: bytes) -> str:
    workbook = load_workbook(filename=io.BytesIO(content), data_only=True)
    rows: List[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            values = [str(cell) for cell in row if cell is not None]
            if values:
                rows.append(" | ".join(values))
    return "\n".join(rows)
